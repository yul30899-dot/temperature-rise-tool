import openpyxl
from openpyxl.chart import LineChart, Reference

wb = openpyxl.Workbook()
ws_data = wb.active
ws_data.title = "RawData"

# Add dummy data
ws_data.append(["Time", "CH1", "CH2"])
ws_data.append(["08:00:00", 25, 26])
ws_data.append(["08:00:10", 30, 35])
ws_data.append(["08:00:20", 40, 50])

ws_chart = wb.create_sheet("Chart")
chart = LineChart()
chart.title = "Temp"
data = Reference(ws_data, min_col=2, min_row=1, max_col=3, max_row=4)
cats = Reference(ws_data, min_col=1, min_row=2, max_row=4)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
ws_chart.add_chart(chart, "A1")

wb.save("test_chart.xlsx")
